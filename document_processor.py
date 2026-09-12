import fitz
from pptx import Presentation


def extract_pdf(file_path):
    """Extract text from every page of a PDF."""
    pages = []

    document = fitz.open(file_path)

    for page_number, page in enumerate(document, start=1):
        text = page.get_text()

        pages.append({
            "page": page_number,
            "text": text
        })

    document.close()
    return pages


def extract_pptx(file_path):
    """Extract text from every slide of a PowerPoint."""
    slides = []

    presentation = Presentation(file_path)

    for slide_number, slide in enumerate(presentation.slides, start=1):
        text_parts = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text_parts.append(shape.text)

        slides.append({
            "page": slide_number,
            "text": "\n".join(text_parts)
        })

    return slides